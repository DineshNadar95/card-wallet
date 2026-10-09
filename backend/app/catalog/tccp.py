"""Import the CFPB Terms of Credit Card Plans (TCCP) survey into the `cards` table.

Source: https://www.consumerfinance.gov/data-research/credit-card-data/terms-credit-card-plans-survey/
TCCP has fees and APRs but no per-category earn rates; those come from the scraper.

Usage (from backend/):
    uv run python -m app.catalog.tccp                 # download the release below
    uv run python -m app.catalog.tccp --file x.xlsx   # use a local copy
"""

import argparse
import re
import sys
import urllib.request
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db.models import Card

TCCP_URL = "https://files.consumerfinance.gov/f/documents/cfpb_tccp-data_2025-12-31.xlsx"

Row = dict[str, Any]

_NETWORKS = [
    (re.compile(r"\bvisa\b", re.I), "Visa"),
    (re.compile(r"\bmaster\s?card\b", re.I), "Mastercard"),
    (re.compile(r"\b(american express|amex)\b", re.I), "American Express"),
    (re.compile(r"\bdiscover\b", re.I), "Discover"),
]


def source_label(as_of: date) -> str:
    half = "Jan–Jun" if as_of.month <= 6 else "Jul–Dec"
    return f"CFPB TCCP {half} {as_of.year}"


def _text(v: Any) -> str | None:
    if v is None:
        return None
    s = " ".join(str(v).split())
    return s or None


def _num(v: Any) -> Decimal | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    try:
        return Decimal(str(v).strip())
    except ArithmeticError:
        return None


def _pct(v: Any) -> Decimal | None:
    """TCCP stores rates as fractions (0.139); we store percent (13.900)."""
    n = _num(v)
    return None if n is None else (n * 100).quantize(Decimal("0.001"))


def _yes(v: Any) -> bool | None:
    s = _text(v)
    return None if s is None else s.lower() == "yes"


def _date(v: Any) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = _text(v)
    return date.fromisoformat(s[:10]) if s else None


def _annual_fee(r: Row) -> Decimal | None:
    kind = (_text(r.get("Periodic Fee Type")) or "").lower()
    if not kind:
        return Decimal("0")  # no periodic fee reported
    if "annual" in kind and (fee := _num(r.get("Annual Fee"))) is not None:
        return fee
    if "monthly" in kind and (fee := _num(r.get("Monthly Fee"))) is not None:
        return (fee * 12).quantize(Decimal("0.01"))
    return None  # a fee exists but the amount isn't a simple number


def _fx_fee(r: Row) -> Decimal | None:
    charged = _yes(r.get("Foreign Transaction Fees?"))
    if charged is False:
        return Decimal("0")
    return _pct(r.get("Foreign Transaction Fee (%)")) if charged else None


def _late_fee(r: Row) -> Decimal | None:
    if _yes(r.get("Late Fees?")) is False:
        return Decimal("0")
    return _num(r.get("Late Fee ($)"))


def _network(name: str) -> str | None:
    return next((label for rx, label in _NETWORKS if rx.search(name)), None)


def _card(r: Row) -> Row | None:
    issuer, name = _text(r.get("Institution Name")), _text(r.get("Product Name"))
    as_of = _date(r.get("Report Date"))
    if not issuer or not name or not as_of:
        return None
    return {
        "issuer": issuer,
        "name": name,
        "network": _network(name),
        "annual_fee": _annual_fee(r),
        "purchase_apr_min": _pct(r.get("Purchase APR min")),
        "purchase_apr_max": _pct(r.get("Purchase APR max")),
        "fx_fee": _fx_fee(r),
        "late_fee": _late_fee(r),
        "cash_advance_apr_min": _pct(r.get("Advance APR min")),
        "cash_advance_apr_max": _pct(r.get("Advance APR max")),
        "secured": _yes(r.get("Secured Card")),
        "credit_tiers": _text(r.get("Targeted Credit Tiers")),
        "rewards_types": _text(r.get("Rewards")),
        "availability": _text(r.get("Availability of Credit Card Plan")),
        "website": _text(r.get("Website for Consumer")),
        "as_of": as_of,
    }


def parse_tccp(path: str | Path) -> tuple[list[Row], int]:
    """Return (cards, skipped). Skips blank rows and repeated (issuer, name) pairs."""
    ws = load_workbook(path, read_only=True, data_only=True).active
    rows = ws.iter_rows(values_only=True)

    # The sheet opens with a title block; the header is the row naming these columns.
    for header in rows:
        names = [_text(c) for c in header]
        if "Institution Name" in names and "Product Name" in names:
            break
    else:
        raise ValueError("TCCP header row not found (no 'Institution Name' / 'Product Name')")

    cards: list[Row] = []
    seen: set[tuple[str, str]] = set()
    skipped = 0
    for values in rows:
        if not any(v is not None for v in values):
            continue
        card = _card({n: v for n, v in zip(names, values, strict=False) if n})
        if card is None or (card["issuer"], card["name"]) in seen:
            skipped += 1
            continue
        seen.add((card["issuer"], card["name"]))
        cards.append(card)
    return cards, skipped


def import_cards(db: Session, rows: list[Row], source_url: str) -> tuple[int, int]:
    """Upsert by (issuer, name). Returns (inserted, updated). Never touches earn rates."""
    existing = {(c.issuer, c.name): c for c in db.scalars(select(Card))}
    inserted = updated = 0
    for row in rows:
        values = {**row, "source": source_label(row["as_of"]), "source_url": source_url}
        card = existing.get((row["issuer"], row["name"]))
        if card is None:
            db.add(Card(**values))
            inserted += 1
        else:
            for k, v in values.items():
                setattr(card, k, v)
            updated += 1
    db.commit()
    return inserted, updated


def main(argv: list[str] | None = None) -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawTextHelpFormatter)
    ap.add_argument("--file", type=Path, help="local .xlsx instead of downloading")
    ap.add_argument("--url", default=TCCP_URL)
    args = ap.parse_args(argv)

    path = args.file
    if path is None:
        path = Path("tccp-download.xlsx")
        print(f"Downloading {args.url} …")
        req = urllib.request.Request(args.url, headers={"User-Agent": "card-wallet-importer"})
        with urllib.request.urlopen(req, timeout=120) as res:
            path.write_bytes(res.read())

    rows, skipped = parse_tccp(path)
    from app.db.session import get_engine

    with Session(get_engine()) as db:
        inserted, updated = import_cards(db, rows, source_url=args.url)
    print(f"{len(rows)} cards: {inserted} new, {updated} updated, {skipped} rows skipped")
    if args.file is None:
        path.unlink()


if __name__ == "__main__":
    sys.exit(main())
